#!/usr/bin/env python3
"""Build the COMAVI manuscript (.docx) from its text source.

WHY THIS EXISTS
---------------
Until v7.14 the manuscript was a hand-edited Word file. Every correction was a
run-level patch, references were renumbered by hand, and the tables were typed
values that nothing recomputed -- which is how a VWF pathogenic count of 0 and
a measured-effect count built on absolute values survived several audits. The
manuscript is now generated:

  docs/manuscript/COMAVI_manuscript_source.txt   the text, with markup
  docs/manuscript/references.json                keyed, CrossRef-verified entries
  docs/manuscript/template.docx                  styles, page setup, line numbers
  reference_outputs/...                          every table value (read here)

Citations are written as keys and numbered by first appearance at build time,
so adding or moving a citation can no longer leave the list out of order.
Table 1 and Table 2 are computed from the committed records on every build.

Run from the repo root:
    python scripts/build_manuscript.py            # write docs/COMAVI_manuscript_current.docx
    python scripts/build_manuscript.py --check    # rebuild in memory, compare text, write nothing
    python scripts/build_manuscript.py --out X.docx
"""
from PIL import Image
import argparse
import copy
import json
import pathlib
import re
import sys
import tempfile

import pandas as pd
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Inches, Pt

REPO = pathlib.Path(__file__).resolve().parents[1]
MS = REPO / "docs" / "manuscript"
SOURCE = MS / "COMAVI_manuscript_source.txt"
REFS = MS / "references.json"
TEMPLATE = MS / "template.docx"
OUT = REPO / "docs" / "COMAVI_manuscript_current.docx"
PROVENANCE = MS / "build_provenance.json"
FIGDIR = REPO / "figures" / "submission"
RO = REPO / "reference_outputs"

TABLE_PT = 9
HEADER_FILL = "D9D9D9"
M_NS = "http://schemas.openxmlformats.org/officeDocument/2006/math"

# ---------------------------------------------------------------- source parse

def parse_source(text):
    """Return a list of blocks: (kind, payload, provenance)."""
    blocks, prov = [], None
    para = []

    def flush():
        nonlocal para, prov
        if para:
            blocks.append(("p", " ".join(para), prov))
            para, prov = [], None

    for raw in text.splitlines():
        line = raw.rstrip()
        if line.startswith("%"):
            m = re.match(r"%from:\s*(.+)", line)
            if m:
                flush()
                prov = m.group(1).strip()
            continue
        if not line.strip():
            flush()
            continue
        if line.startswith("#"):
            flush()
            level = len(line) - len(line.lstrip("#"))
            blocks.append(("h%d" % level, line.lstrip("#").strip(), None))
            continue
        if line.startswith("@"):
            flush()
            name, _, rest = line[1:].partition(" ")
            blocks.append(("@" + name, rest.strip(), prov))
            prov = None
            continue
        para.append(line.strip())
    flush()
    return blocks


# ---------------------------------------------------------------- citations

CITE = re.compile(r"\[(@[^\]]+)\]")


def cite_keys(s):
    return re.findall(r"@([A-Za-z0-9_]+)", s)


def number_citations(blocks):
    """First-appearance order over everything a reader meets before References."""
    order = []
    for kind, payload, _ in blocks:
        if kind == "h1" and payload == "References":
            break
        if kind in ("p", "@figure", "@table"):
            for m in CITE.finditer(payload):
                for k in cite_keys(m.group(1)):
                    if k not in order:
                        order.append(k)
    return {k: i + 1 for i, k in enumerate(order)}


def compress(nums):
    nums = sorted(set(nums))
    out, i = [], 0
    while i < len(nums):
        j = i
        while j + 1 < len(nums) and nums[j + 1] == nums[j] + 1:
            j += 1
        if j == i:
            out.append(str(nums[i]))
        elif j == i + 1:
            out += [str(nums[i]), str(nums[j])]
        else:
            out.append("%d\u2013%d" % (nums[i], nums[j]))
        i = j + 1
    return ",".join(out)


def render_citations(s, numbers):
    def rep(m):
        keys = cite_keys(m.group(1))
        missing = [k for k in keys if k not in numbers]
        if missing:
            raise KeyError("uncited or unknown reference keys: %s" % missing)
        return "[" + compress([numbers[k] for k in keys]) + "]"
    return CITE.sub(rep, s)


# ---------------------------------------------------------------- inline runs

INLINE = re.compile(r"\{(i|b|sup|sub)\|([^}]*)\}")


def add_runs(par, text, bold=False, italic=False, size=None):
    pos = 0
    for m in INLINE.finditer(text):
        if m.start() > pos:
            _run(par, text[pos:m.start()], bold, italic, size)
        kind, body = m.group(1), m.group(2)
        r = _run(par, body, bold or kind == "b", italic or kind == "i", size)
        if kind == "sup":
            r.font.superscript = True
        if kind == "sub":
            r.font.subscript = True
        pos = m.end()
    if pos < len(text):
        _run(par, text[pos:], bold, italic, size)


def _run(par, text, bold, italic, size):
    r = par.add_run(text)
    if bold:
        r.bold = True
    if italic:
        r.italic = True
    if size:
        r.font.size = Pt(size)
    return r


def plain(text):
    return INLINE.sub(lambda m: m.group(2), text)


# ---------------------------------------------------------------- equations

def _mr(text, plain_style=False):
    sty = '<m:rPr><m:sty m:val="p"/></m:rPr>' if plain_style else ""
    t = text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
    return '<m:r>%s<m:t xml:space="preserve">%s</m:t></m:r>' % (sty, t)


def _group(s, i):
    """Read a {...} group starting at s[i] == '{'; return (content, next index)."""
    assert s[i] == "{", s[i:]
    depth, j = 0, i
    while j < len(s):
        if s[j] == "{":
            depth += 1
        elif s[j] == "}":
            depth -= 1
            if depth == 0:
                return s[i + 1:j], j + 1
        j += 1
    raise ValueError("unbalanced braces in equation: %s" % s)


def omml(s):
    """Tiny LaTeX subset -> OMML: \\frac{}{}, \\max_{}, X_{}, plain text."""
    out, i, buf = [], 0, ""

    def flush_buf():
        nonlocal buf
        if buf:
            out.append(_mr(buf))
            buf = ""

    while i < len(s):
        if s.startswith("\\frac", i):
            flush_buf()
            num, i = _group(s, i + 5)
            den, i = _group(s, i)
            out.append("<m:f><m:num>%s</m:num><m:den>%s</m:den></m:f>" % (omml(num), omml(den)))
        elif s.startswith("\\max_", i):
            flush_buf()
            lim, i = _group(s, i + 5)
            out.append('<m:limLow><m:e>%s</m:e><m:lim>%s</m:lim></m:limLow>'
                       % (_mr("max", plain_style=True), omml(lim)))
        elif s[i] == "_" and i + 1 < len(s) and s[i + 1] == "{":
            base = buf[-1] if buf else ""
            buf = buf[:-1]
            if base == "G" and buf.endswith("ΔΔ"):
                base, buf = "ΔΔG", buf[:-2]
            flush_buf()
            sub, i = _group(s, i + 1)
            out.append("<m:sSub><m:e>%s</m:e><m:sub>%s</m:sub></m:sSub>" % (_mr(base), omml(sub)))
        else:
            buf += s[i]
            i += 1
    flush_buf()
    return "".join(out)


def add_equation(doc, latex):
    p = doc.add_paragraph()
    xml = ('<m:oMathPara xmlns:m="%s"><m:oMath>%s</m:oMath></m:oMathPara>' % (M_NS, omml(latex)))
    p._p.append(parse_xml(xml))
    return p


# ---------------------------------------------------------------- tables

def _shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    tcPr.append(parse_xml('<w:shd %s w:val="clear" w:color="auto" w:fill="%s"/>' % (nsdecls("w"), fill)))


def _grid(table):
    tblPr = table._tbl.tblPr
    for el in tblPr.findall(qn("w:tblBorders")):
        tblPr.remove(el)
    borders = "".join('<w:%s w:val="single" w:sz="4" w:space="0" w:color="000000"/>' % b
                      for b in ("top", "left", "bottom", "right", "insideH", "insideV"))
    tblPr.append(parse_xml("<w:tblBorders %s>%s</w:tblBorders>" % (nsdecls("w"), borders)))


def write_table(doc, header, rows, italic_cols=(), widths=None, left_cols=()):
    t = doc.add_table(rows=1 + len(rows), cols=len(header))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    _grid(t)
    for ri, values in enumerate([header] + rows):
        for ci, v in enumerate(values):
            cell = t.rows[ri].cells[ci]
            cell.text = ""
            par = cell.paragraphs[0]
            par.alignment = WD_ALIGN_PARAGRAPH.LEFT if ci in left_cols else WD_ALIGN_PARAGRAPH.CENTER
            pf = par.paragraph_format
            pf.space_before, pf.space_after, pf.line_spacing = Pt(1), Pt(1), 1.0
            add_runs(par, str(v), bold=(ri == 0), italic=(ri > 0 and ci in italic_cols), size=TABLE_PT)
            if ri == 0:
                _shade(cell, HEADER_FILL)
            if widths:
                cell.width = Inches(widths[ci])
    # header row repeats on each page
    trPr = t.rows[0]._tr.get_or_add_trPr()
    trPr.append(parse_xml("<w:tblHeader %s/>" % nsdecls("w")))
    return t


SYSTEMS = [  # display name, system id, gene, structure
    ("BRCA1 BRCT tandem domain", "brca1_brct", "BRCA1", "PDB 1JNX"),
    ("BRCA1–BARD1 RING heterodimer", "brca1_bard1", "BRCA1", "PDB 1JM7"),
    ("Calmodulin–CaV1.2 IQ domain", "cam_cav12", "CALM1", "AlphaFold 3 model"),
    ("Complement factor H–C3b", "cfh_c3b", "CFH", "PDB 2WII"),
    ("Hemoglobin α–β dimer", "hemoglobin_dimer", "HBB", "PDB 2HHB"),
    ("Hemoglobin tetramer", "hemoglobin_tetramer", "HBB", "PDB 2HHB"),
    ("KRAS–RAF1 RAS-binding domain", "kras_craf", "KRAS", "PDB 6XI7"),
    ("MLH1–PMS2 C-terminal heterodimer", "mlh1_pms2", "MLH1/PMS2", "AlphaFold 3 model"),
    ("MSH2–MSH6 heterodimer", "msh2_msh6", "MSH2", "AlphaFold 3 model"),
    ("PI3K p110α–p85α", "pi3k", "PIK3CA/PIK3R1", "AlphaFold 3 model"),
    ("SMAD4–SMAD3 MH2 complex", "smad4_smad3", "SMAD4", "AlphaFold 3 model"),
    ("Troponin I–troponin C", "troponin_ic", "TNNI3", "AlphaFold 3 model"),
    ("VHL–elongin C", "vhl_elonginc", "VHL", "AlphaFold 3 model"),
    ("VWF A1–GPIbα", "vwf_gpiba", "VWF", "PDB 1SQ0"),
]


def table1_rows():
    can = pd.read_csv(RO / "scored_61var_canonical.csv")
    assert set(can.system) == {s[1] for s in SYSTEMS}, sorted(set(can.system) ^ {s[1] for s in SYSTEMS})
    rows, tot = [], [0, 0, 0]
    for name, sid, gene, struct in SYSTEMS:
        sub = can[can.system == sid]
        n, graded = len(sub), int(sub.mech_consistency_t25.notna().sum())
        ph = sub.phenotype.astype(str).str.lower()
        labelled = ph.str.startswith(("pathogenic", "benign", "likely"))
        path = int(ph.str.startswith("pathogenic").sum())
        rows.append([name, gene, n, graded, path if labelled.any() else "–", struct])
        tot[0] += n
        tot[1] += graded
        tot[2] += path
    assert tot[:2] == [61, 57], tot
    rows.append(["Total", "", tot[0], tot[1], tot[2], ""])
    return ["System", "Gene", "Variants", "Graded", "Pathogenic", "Structure"], rows


THRESH_LABELS = ["1.0", "1.5", "2.0", "2.5", "Bound"]


def table2_rows():
    sw = json.loads((RO / "COMAVI_threshold_sweep.json").read_text())["rows"]
    assert [r["threshold_tag"] for r in sw] == ["t10", "t15", "t20", "t25", "tSAP"]
    can = pd.read_csv(RO / "scored_61var_canonical.csv")
    g = can[can.mech_consistency_t25.notna()]
    w = {"consistent": 1.0, "partial": 0.5, "inconsistent": 0.0}
    struct = g.expected_mech_class != "structurally_silent"
    me = json.loads((RO / "COMAVI_measured_effect_tables.json").read_text())["measured_effects_recovered"]
    pm = pd.read_csv(RO / "isds_v1" / "ISDS_v1_primary_metrics.csv").set_index("score")
    f3 = lambda x: "%.3f" % x

    def arm(mask):
        out = []
        for tag in ["t10", "t15", "t20", "t25", "tSAP"]:
            out.append(f3(g.loc[mask, "mech_consistency_%s" % tag].map(w).mean()))
        return out

    da = lambda r: (r["axis_monomer_k"] + r["axis_fold_k"] + r["axis_binding_k"]) / \
                   (r["axis_monomer_n"] + r["axis_fold_n"] + r["axis_binding_n"])
    n_axes = sw[0]["axis_monomer_n"] + sw[0]["axis_fold_n"] + sw[0]["axis_binding_n"]
    auc = f3(pm.loc["isds_v1", "roc_auc"])
    rows = [
        ["Mechanism-pattern score", "%d graded variants" % len(g), "Mean whole-variant grade"] + [f3(r["whole_variant"]) for r in sw],
        ["Structural arm", "%d structural-mechanism variants" % int(struct.sum()), "Mean grade of variants expected to carry a lesion"] + arm(struct),
        ["No-lesion arm", "%d no-lesion variants" % int((~struct).sum()), "Mean grade of variants expected to be silent"] + arm(~struct),
        ["Detection", "%d structural-mechanism variants" % sw[0]["detection_n"], "Any axis reaches the threshold"] + [f3(r["detection"]) for r in sw],
        ["Attribution", "%d structural-mechanism variants" % sw[0]["attribution_n"], "The axis the literature implicates reaches the threshold"] + [f3(r["attribution"]) for r in sw],
        ["Correct rejection", "%d no-lesion variants" % sw[0]["correct_rejection_n"], "No axis reaches the threshold"] + [f3(r["correct_rejection"]) for r in sw],
        ["Direction agreement, energy axes", "%d axis expectations" % n_axes, "Call matches the committed direction"] + [f3(da(r)) for r in sw],
        ["Direction agreement, tier", "%d variants" % sw[0]["axis_tier_n"], "Tier 1–2 matches the committed structural class"] + [f3(r["axis_tier"]) for r in sw],
        ["Measured destabilizations recovered", "%d measured comparisons" % me["population"]["n"], "Measured ≥ 1.0 kcal/mol; predicted ≥ threshold with the same sign"] + [f3(r["fraction"]) for r in me["by_threshold"]],
        ["Prioritization ROC AUC", "%d prioritization variants" % int(pm.loc["isds_v1", "n"]), "A structural mechanism outranks a no-lesion variant"] + [auc] * 5,
    ]
    return ["Metric", "Population", "What it counts"] + THRESH_LABELS, rows


def table3_rows():
    header = ["Condition", "Grade", "Score"]
    rows = [
        ["No-lesion variant, and no modeled energy effect at the threshold", "Consistent", "1.0"],
        ["All expected effects recovered in the correct direction, with no unsupported effect", "Consistent", "1.0"],
        ["All expected effects recovered, but an axis expected to be neutral also reaches the threshold", "Partial", "0.5"],
        ["One of several expected effects missed, without an unsupported effect", "Partial", "0.5"],
        ["Predicted change opposite in direction to the expected effect", "Inconsistent", "0.0"],
        ["No-lesion variant with an energy call; expected effect missed entirely; or a missed effect with a wrong call", "Inconsistent", "0.0"],
        ["Structure unusable, required context absent, mechanism not established or no applicable axis", "Not graded", "–"],
    ]
    return header, rows


TABLES = {
    "table1": (table1_rows, dict(italic_cols=(1,), left_cols=(0, 5), widths=(2.1, 1.0, 0.7, 0.6, 0.8, 1.3))),
    "table2": (table2_rows, dict(left_cols=(0, 1, 2), widths=(1.35, 1.1, 1.6, 0.48, 0.48, 0.48, 0.48, 0.5))),
    "table3": (table3_rows, dict(left_cols=(0,), widths=(4.6, 1.0, 0.7))),
}


# ---------------------------------------------------------------- document

def clear_body(doc):
    body = doc.element.body
    for el in list(body):
        if el.tag != qn("w:sectPr"):
            body.remove(el)


def caption(doc, label, title, legend, numbers):
    p = doc.add_paragraph()
    add_runs(p, label + ". " + render_citations(title, numbers), bold=True)
    if legend:
        p.add_run(" ")
        add_runs(p, render_citations(legend, numbers))
    return p


def build(out_path):
    blocks = parse_source(SOURCE.read_text())
    refs = json.loads(REFS.read_text())
    numbers = number_citations(blocks)
    unknown = [k for k in numbers if k not in refs]
    if unknown:
        raise KeyError("citation keys missing from references.json: %s" % unknown)
    uncited = [k for k in refs if k not in numbers]
    if uncited:
        raise ValueError("references never cited before the reference list: %s" % uncited)

    doc = Document(str(TEMPLATE))
    clear_body(doc)
    prov = []
    for kind, payload, src in blocks:
        if kind == "@title":
            p = doc.add_paragraph(style="Title")
            add_runs(p, payload)
        elif kind == "@short_title":
            p = doc.add_paragraph()
            add_runs(p, "Short title: ", bold=True)
            add_runs(p, payload)
        elif kind in ("@authors", "@affiliations", "@corresponding"):
            add_runs(doc.add_paragraph(), payload)
        elif kind in ("h1", "h2", "h3"):
            doc.add_paragraph(payload, style="Heading %s" % kind[1])
        elif kind == "p":
            p = doc.add_paragraph()
            add_runs(p, render_citations(payload, numbers))
            prov.append({"text": plain(render_citations(payload, numbers)), "from": src})
        elif kind == "@figure":
            label, title, legend = [x.strip() for x in payload.split("|", 2)]
            img = FIGDIR / (label.replace(" ", "") + ".png")
            if img.exists():
                ip = doc.add_paragraph()
                ip.alignment = WD_ALIGN_PARAGRAPH.CENTER
                # Insert at the printed width the figure was drawn for, so the page
                # shows it at 100% and every label prints at the size it was checked at.
                with Image.open(img) as _im:
                    _w_in = _im.size[0] / float(_im.info.get("dpi", (300, 300))[0])
                ip.add_run().add_picture(str(img), width=Inches(_w_in))
            caption(doc, label, title, legend, numbers)
            prov.append({"text": "%s. %s %s" % (label, plain(title), plain(render_citations(legend, numbers))), "from": "caption"})
        elif kind == "@table":
            tid, label, title, legend = [x.strip() for x in payload.split("|", 3)]
            fn, opts = TABLES[tid]
            header, rows = fn()
            caption(doc, label, title, "", numbers)
            write_table(doc, header, rows, **opts)
            if legend:
                add_runs(doc.add_paragraph(), render_citations(legend, numbers))
            prov.append({"text": "%s. %s %s" % (label, plain(title), plain(render_citations(legend, numbers))), "from": "caption"})
        elif kind == "@eq":
            add_equation(doc, payload)
        elif kind == "@references":
            for k, n in sorted(numbers.items(), key=lambda kv: kv[1]):
                p = doc.add_paragraph()
                add_runs(p, "%d. %s" % (n, refs[k]["text"]))
        elif kind == "@si":
            label, title, legend = [x.strip() for x in payload.split("|", 2)]
            p = doc.add_paragraph()
            add_runs(p, label + ". " + title, bold=True)
            if legend:
                p.add_run(" ")
                add_runs(p, legend)
        else:
            raise ValueError("unknown directive %s" % kind)
    doc.save(str(out_path))
    return {"numbers": numbers, "provenance": prov}


def doc_text(path):
    d = Document(str(path))
    out = []
    for el in d.element.body:
        if el.tag == qn("w:p"):
            out.append("".join(t.text or "" for t in el.iter(qn("w:t"))) +
                       "".join(t.text or "" for t in el.iter("{%s}t" % M_NS)))
        elif el.tag == qn("w:tbl"):
            for tr in el.iter(qn("w:tr")):
                out.append(" | ".join("".join(t.text or "" for t in tc.iter(qn("w:t"))) for tc in tr.iter(qn("w:tc"))))
    return out


def main():
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("--check", action="store_true", help="rebuild and compare text with the committed docx; write nothing")
    ap.add_argument("--out", type=pathlib.Path, default=OUT)
    args = ap.parse_args()
    if args.check:
        with tempfile.TemporaryDirectory() as td:
            tmp = pathlib.Path(td) / "build.docx"
            build(tmp)
            fresh, committed = doc_text(tmp), doc_text(OUT)
        if fresh != committed:
            diff = [i for i, (a, b) in enumerate(zip(fresh, committed)) if a != b]
            print("FAIL: %s differs from a fresh build (%d vs %d blocks; first difference at block %s)"
                  % (OUT.name, len(committed), len(fresh), diff[0] if diff else min(len(fresh), len(committed))))
            return 1
        print("PASS: %s reproduces from its source (%d blocks)" % (OUT.name, len(fresh)))
        return 0
    res = build(args.out)
    PROVENANCE.write_text(json.dumps(res, ensure_ascii=False, indent=1) + "\n")
    print("wrote %s: %d references, %d paragraphs" % (args.out.relative_to(REPO) if args.out.is_relative_to(REPO) else args.out,
                                                      len(res["numbers"]), len(res["provenance"])))
    return 0


if __name__ == "__main__":
    sys.exit(main())
